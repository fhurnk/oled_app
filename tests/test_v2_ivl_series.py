import json
import tempfile
import time
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from oled_app.measurements.ivl import run_ivl_cycle
from oled_app.settings import load_app_settings
from oled_v2.ivl import IvlController, validate_params
from oled_v2.series_service import SeriesService, SeriesConflictError, SeriesNotFoundError
from test_v2_series_service import series_payload
from test_v2_ivl import wait_terminal


class IvlSeriesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.service = SeriesService(self.root)
        active = self.service.create_series(series_payload(self.root))["active"]
        self.target = {"series_path": active["path"], "pixel_id": active["pixels"][0]["pixel_id"]}
        self.controller = IvlController(self.root / "runs", self.service)
        self.input = {"target": self.target, "sweep_end": 0.2, "sweep_increment": 0.1}

    def tearDown(self):
        self.controller.shutdown()
        self.temp.cleanup()

    def force_statuses(self, statuses):
        remaining = iter(statuses)
        def measure(*args, **kwargs):
            result = run_ivl_cycle(*args, **kwargs)
            result["status"] = next(remaining)
            result["opening_voltage"] = None
            return result
        return patch("oled_v2.ivl.run_ivl_cycle", side_effect=measure)

    def wait_decision(self):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            state = self.controller.snapshot()
            if state.get("decision"):
                return state
            if not state["active"]:
                self.fail(str(state))
            time.sleep(0.01)
        self.fail("Opening decision did not arrive")

    def test_preflight_resolves_calibration_without_measurement_files(self):
        result = self.controller.preflight(self.input)
        self.assertEqual(result["target"], self.target)
        self.assertEqual(result["cycles"], 1)
        self.assertFalse(list(Path(self.target["series_path"]).rglob("IVL_*.xlsx")))
        self.assertFalse(self.controller.output_root.exists())
        self.assertEqual(result["luminance_coefficient"],
                         self.service._active.rgb_luminance_coefficient_for_pixel(
                             self.target["pixel_id"], load_app_settings()))

    def test_series_result_persists_history_relative_paths_thumbnail_and_simulator_marker(self):
        self.controller.start(self.input)
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertTrue(state["result"]["journaled"])
        active = self.service.state()["active"]
        pixel = next(p for p in active["pixels"] if p["pixel_id"] == self.target["pixel_id"])
        self.assertTrue(pixel["thumbnail_available"])
        self.assertFalse(Path(pixel["last_ivl_file"]).is_absolute())
        self.assertEqual(active["metrics"]["ivl"], 1)
        self.assertIn("ЭМУЛЯТОР v2", active["history"][0]["notes"])
        workbook = load_workbook(Path(active["path"]) / "series_journal.xlsx", read_only=True)
        try:
            sheet = workbook["Measurements"]
            headers = [cell.value for cell in sheet[1]]
            row = dict(zip(headers, [cell.value for cell in sheet[2]]))
            params = next(json.loads(value) for value in row.values()
                          if isinstance(value, str) and value.startswith('{"com_port"'))
            self.assertEqual(params["hardware_mode"], "simulator")
        finally:
            workbook.close()
        self.service.close_series()
        reopened = self.service.open_series(self.target["series_path"])["active"]
        self.assertEqual(reopened["metrics"]["ivl"], 1)

    def test_stale_target_and_unknown_pixel_rejected(self):
        with self.assertRaises(SeriesConflictError):
            self.controller.start({**self.input, "target": {**self.target, "series_path": "wrong"}})
        with self.assertRaises(SeriesNotFoundError):
            self.controller.start({**self.input, "target": {**self.target, "pixel_id": "../unknown"}})
        self.assertFalse(self.controller.output_root.exists())

    def test_burnout_confirmation_uses_final_status_not_first_cycle(self):
        with self.force_statuses(["BURNED", "NONWORKING"]):
            self.controller.start(self.input)
            state = wait_terminal(self.controller)
        self.assertEqual(state["result"]["cycles"], 2)
        self.assertEqual(state["result"]["status"], "BURNED")
        self.assertTrue(state["result"]["journaled"])

    def test_multiple_cycles_are_saved_in_one_workbook(self):
        with self.force_statuses(["WORKING", "NONWORKING"]):
            self.controller.start({**self.input, "num_cycles": 2, "delay_between_cycles": 0.05})
            state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed")
        self.assertEqual(state["cycle"], 2)
        self.assertEqual(state["point_count"], 6)
        wb = load_workbook(state["result"]["file"], read_only=True)
        try:
            self.assertEqual(wb.sheetnames, ["Summary", "Cycle_1", "Cycle_2"])
        finally:
            wb.close()

    def test_missing_opening_waits_with_outputs_off_and_accepts_one_decision(self):
        with self.force_statuses(["WORKING"]):
            self.controller.start(self.input)
            state = self.wait_decision()
            self.assertTrue(state["safe_shutdown_confirmed"])
            self.assertEqual(self.service.state()["active"]["metrics"]["ivl"], 0)
            decision = {"run_id": state["run_id"], "decision_id": state["decision"]["id"], "value": 3.2}
            with self.assertRaises(RuntimeError):
                self.controller.decide_opening({**decision, "run_id": "old"})
            with self.assertRaises(ValueError):
                self.controller.decide_opening({**decision, "value": float("nan")})
            self.controller.decide_opening(decision)
            with self.assertRaises(RuntimeError):
                self.controller.decide_opening(decision)
            state = wait_terminal(self.controller)
        self.assertTrue(state["result"]["journaled"])
        self.assertEqual(self.service._active.journal.get_pixel(self.target["pixel_id"])["Opening voltage (V)"], 3.2)

    def test_cancel_pending_decision_keeps_workbook_without_journal(self):
        with self.force_statuses(["WORKING"]):
            self.controller.start(self.input)
            self.wait_decision()
            self.controller.shutdown()
        state = self.controller.snapshot()
        self.assertEqual(state["status"], "stopped")
        self.assertTrue(Path(state["result"]["file"]).is_file())
        self.assertFalse(state["result"]["journaled"])
        self.assertEqual(self.service.state()["active"]["metrics"]["ivl"], 0)

    def test_journal_failure_preserves_result_and_releases_operation(self):
        with patch.object(self.service, "record_ivl", side_effect=PermissionError("Workbook locked")):
            self.controller.start(self.input)
            state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "failed")
        self.assertTrue(Path(state["result"]["file"]).is_file())
        self.assertFalse(state["result"]["journaled"])
        self.assertTrue(state["safe_shutdown_confirmed"])

    def test_skip_opening_saves_one_result_without_invented_voltage(self):
        with self.force_statuses(["WORKING"]):
            self.controller.start(self.input)
            state = self.wait_decision()
            self.controller.decide_opening({"run_id": state["run_id"],
                "decision_id": state["decision"]["id"], "value": None})
            state = wait_terminal(self.controller)
        self.assertTrue(state["result"]["journaled"])
        self.assertIsNone(state["result"]["opening_voltage"])

    def test_unconfirmed_shutdown_never_journals_success(self):
        with patch("oled_v2.ivl.safe_shutdown_smu", return_value=False):
            self.controller.start(self.input)
            state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "failed")
        self.assertFalse(state["safe_shutdown_confirmed"])
        self.assertEqual(self.service.state()["active"]["metrics"]["ivl"], 0)

    def test_repeat_preserves_previous_workbook_and_history(self):
        self.controller.start(self.input)
        first = wait_terminal(self.controller)
        self.controller.start(self.input)
        second = wait_terminal(self.controller)
        self.assertNotEqual(first["result"]["file"], second["result"]["file"])
        self.assertTrue(Path(first["result"]["file"]).is_file())
        self.assertTrue(Path(second["result"]["file"]).is_file())
        self.assertEqual(len(self.service.state()["active"]["history"]), 2)

    def test_cycles_validation_and_fractional_delay(self):
        self.assertEqual(validate_params({"delay_between_cycles": 0.15}).delay_between_cycles, 0.15)
        for payload in ({"num_cycles": 1.5}, {"num_cycles": 0}, {"burned_confirmation_cycles": 1.5}):
            with self.assertRaises(ValueError):
                validate_params(payload)


if __name__ == "__main__":
    unittest.main()
