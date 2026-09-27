from __future__ import annotations

import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from oled_app.hardware import uninstall_simulator_modules
from oled_app.measurements.ivl import IVLParams
from oled_app.settings import load_app_settings
from oled_v2.series_service import SeriesService, SeriesValidationError
from oled_v2.stability import StabilityController, validate_params


FAST_VOLTAGE = {
    "control_mode": "voltage",
    "voltage_setpoint_V": 1.0,
    "voltage_start": 1.0,
    "voltage_limit": 5.0,
    "current_limit_mA": 10.0,
    "measurement_time_s": 0.12,
    "sample_interval_s": 0.01,
    "autosave_interval_s": 60.0,
}


def wait_terminal(controller: StabilityController, timeout: float = 15.0):
    deadline = time.monotonic() + timeout
    while controller.snapshot()["active"] and time.monotonic() < deadline:
        time.sleep(0.01)
    state = controller.snapshot()
    if state["active"]:
        raise AssertionError("Stability controller did not reach a terminal state")
    return state


class StabilityControllerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.controller = StabilityController(self.root / "standalone")

    def tearDown(self):
        self.controller.shutdown()
        uninstall_simulator_modules()
        self.temp.cleanup()

    def test_validation_rejects_conflicting_limits(self):
        with self.assertRaisesRegex(ValueError, "Уставка тока"):
            validate_params({"current_setpoint_mA": 2.0, "current_limit_mA": 1.0})
        with self.assertRaisesRegex(ValueError, "Интервал точки"):
            validate_params({"measurement_time_s": 1.0, "sample_interval_s": 2.0})

    def test_standalone_voltage_run_updates_setpoint_and_builds_workbook(self):
        checked = self.controller.preflight(FAST_VOLTAGE)
        self.assertEqual(checked["effective_voltage_start"], 1.0)
        started = self.controller.start(FAST_VOLTAGE)
        deadline = time.monotonic() + 5
        while self.controller.snapshot()["point_count"] < 1 and time.monotonic() < deadline:
            time.sleep(0.01)
        changed = self.controller.setpoint({"run_id": started["run_id"], "value": 1.5})
        self.assertEqual(changed["current_setpoint"], 1.5)
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertTrue(state["safe_shutdown_confirmed"])
        self.assertEqual(state["result"]["final_setpoint"], 1.5)
        self.assertFalse(state["result"]["journaled"])
        self.assertTrue(Path(state["result"]["file"]).is_file())
        workbook = load_workbook(state["result"]["file"], read_only=True)
        try:
            self.assertIn("Data", workbook.sheetnames)
        finally:
            workbook.close()

    def test_stop_keeps_completed_points_and_confirms_shutdown(self):
        self.controller.start({**FAST_VOLTAGE, "measurement_time_s": 2.0})
        deadline = time.monotonic() + 5
        while self.controller.snapshot()["point_count"] < 1 and time.monotonic() < deadline:
            time.sleep(0.01)
        self.controller.stop()
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "stopped", state)
        self.assertTrue(state["safe_shutdown_confirmed"])
        self.assertTrue(state["result"]["stopped_by_user"])
        self.assertTrue(Path(state["result"]["file"]).is_file())

    def test_series_current_mode_uses_ivl_start_and_writes_journal(self):
        service = SeriesService(self.root / "series")
        active = service.create_series({
            "root": str(self.root / "series"), "deposition_date": "2026-09-28",
            "keyword": "stability", "series_led_color": "green",
            "quarter_bases": {str(number): "Q" for number in range(1, 5)},
            "quarter_descriptions": {str(number): "Simulator" for number in range(1, 5)},
        })["active"]
        pixel_id = active["pixels"][0]["pixel_id"]
        self.controller.series_service = service
        with self.assertRaises(SeriesValidationError):
            self.controller.preflight({
                **FAST_VOLTAGE,
                "target": {"series_path": active["path"], "pixel_id": pixel_id},
            })
        ivl_file = Path(active["path"]) / "seed_ivl.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cycle_1"
        sheet.append(["Voltage OLED / LED measured (V)", "Current OLED / LED (mA)"])
        sheet.append([2.0, 0.0])
        sheet.append([4.0, 1.0])
        workbook.save(ivl_file)
        target = service.ivl_target(
            {"series_path": active["path"], "pixel_id": pixel_id},
            IVLParams(), load_app_settings(),
        )
        service.record_ivl(target, IVLParams(), {
            "status": "WORKING", "file": str(ivl_file), "run_id": "seed",
            "ivl_diagnosis": "seed", "opening_voltage": 2.0,
            "max_current_mA": 1.0, "max_photo_uA": 1.0,
        })
        payload = {
            "control_mode": "current", "current_setpoint_mA": 0.5,
            "voltage_start": 1.0, "voltage_limit": 5.0, "current_limit_mA": 10.0,
            "measurement_time_s": 0.08, "sample_interval_s": 0.01,
            "autosave_interval_s": 60.0,
            "target": {"series_path": active["path"], "pixel_id": pixel_id},
            "use_ivl_start_voltage": True,
        }
        checked = self.controller.preflight(payload)
        self.assertAlmostEqual(checked["ivl_voltage_at_target"], 3.0)
        self.assertAlmostEqual(checked["effective_voltage_start"], 2.7)
        self.controller.start(payload)
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertTrue(state["result"]["journaled"])
        reopened = service.open_series(active["path"])["active"]
        pixel = next(item for item in reopened["pixels"] if item["pixel_id"] == pixel_id)
        self.assertTrue(pixel["last_stability_file"])
        self.assertIn("ЭМУЛЯТОР v2", reopened["history"][-1]["notes"])

    def test_stale_setpoint_is_rejected(self):
        state = self.controller.start(FAST_VOLTAGE)
        wait_terminal(self.controller)
        with self.assertRaises(RuntimeError):
            self.controller.setpoint({"run_id": state["run_id"], "value": 2.0})


if __name__ == "__main__":
    unittest.main()
