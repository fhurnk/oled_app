from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import tempfile
import time
import unittest

import numpy as np
from openpyxl import load_workbook

from oled_app.measurements.ivl import IVLParams
from oled_app.measurements.spectrum import SpectrumParams
from oled_app.processing.ivl_results import save_ivl_workbook
from oled_app.processing.spectrum_results import create_spectrum_workbook
from oled_app.series.manager import SeriesManager
from oled_app.settings import DEFAULT_APP_SETTINGS
from oled_v2.recalculation import RecalculationService, RecalculationValidationError
from oled_v2.series_service import SeriesService


class V2RecalculationServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.series_service = SeriesService(root / "series")
        active = self.series_service.create_series({
            "root": str(root / "series"),
            "deposition_date": "2026-09-30",
            "keyword": "recalculation-test",
            "series_led_color": "green",
            "quarter_bases": {str(number): "Q" for number in range(1, 5)},
            "quarter_descriptions": {str(number): f"Group {number}" for number in range(1, 5)},
        })["active"]
        self.series_path = Path(active["path"])
        self.pixel_id = active["pixels"][0]["pixel_id"]
        self.settings = deepcopy(DEFAULT_APP_SETTINGS)

        def load_settings():
            return deepcopy(self.settings)

        def save_settings(value):
            self.settings = deepcopy(value)

        self.service = RecalculationService(
            settings_loader=load_settings,
            settings_saver=save_settings,
        )

    def tearDown(self) -> None:
        self.service.shutdown()
        self.temp.cleanup()

    def _wait(self) -> dict:
        deadline = time.monotonic() + 12.0
        while time.monotonic() < deadline:
            state = self.service.snapshot()
            if not state["active"]:
                return state
            time.sleep(0.02)
        self.fail("Recalculation did not finish in time")

    def _seed_spectrum(self) -> Path:
        source = self.series_path / "spectrum_source.xlsx"
        params = SpectrumParams(voltage_start=2.0, voltage_end=3.0, voltage_step=1.0)
        workbook = create_spectrum_workbook(source, self.pixel_id, params, [2.0, 3.0])
        summary = workbook["Сводка"]
        processed = workbook["Processed counts per s"]
        wavelengths = np.linspace(400.0, 700.0, 301)
        for point, (voltage, photo, scale) in enumerate(
            ((2.0, 2.0, 1.0), (3.0, 4.0, 1.8)), start=1
        ):
            row = 21 + point
            summary.cell(row, 1, point)
            summary.cell(row, 2, voltage)
            summary.cell(row, 7, photo)
            summary.cell(row, 8, photo * 5.0)
            summary.cell(row, 13, "GOOD")
            column = point + 1
            processed.cell(2, column, point)
            processed.cell(3, column, voltage)
            spectrum = scale * 1000.0 * np.exp(-0.5 * ((wavelengths - 530.0) / 35.0) ** 2)
            for data_row, (wavelength, intensity) in enumerate(zip(wavelengths, spectrum), start=21):
                processed.cell(data_row, 1, float(wavelength))
                processed.cell(data_row, column, float(intensity))
        workbook.save(source)
        workbook.close()
        manager = SeriesManager(self.series_path)
        manager.journal.update_after_measurement(
            "SPECTRUM", self.pixel_id, "GOOD", source, {}, opening_voltage=2.0
        )
        return source

    def _seed_ivl(self) -> Path:
        output = self.series_path / "ivl_source.xlsx"
        cycle = {
            "cycle": 1,
            "status": "WORKING",
            "status_desc": "ok",
            "current_limit_reached": False,
            "max_photo_uA": 2.0,
            "max_current_mA": 1.0,
            "opening_voltage": 2.0,
            "data": [{
                "Point": 1,
                "Voltage set (V)": 2.0,
                "Voltage OLED / LED measured (V)": 1.99,
                "Current OLED / LED (mA)": 1.0,
                "Current density (mA/cm^2)": 100.0,
                "Voltage photodiode measured (V)": -5.0,
                "Photodiode current (uA)": 2.0,
                "Luminance (cd/m^2)": 2.0,
                "Measurement time (s)": 0.1,
            }],
        }
        save_ivl_workbook(self.pixel_id, output, IVLParams(), [cycle])
        manager = SeriesManager(self.series_path)
        manager.journal.update_after_measurement(
            "IVL", self.pixel_id, "WORKING", output, {},
            opening_voltage=2.0, max_current_mA=1.0, max_photo_uA=2.0,
        )
        return output

    def test_options_expose_groups_thresholds_and_counts(self) -> None:
        self._seed_spectrum()
        self._seed_ivl()

        options = self.service.options(self.series_path)

        self.assertEqual(len(options["groups"]), 4)
        candidate_group = next(group for group in options["groups"] if group["candidates"])
        self.assertEqual(candidate_group["candidates"][0]["pixel_id"], self.pixel_id)
        self.assertEqual(options["measurement_counts"]["SPECTRUM"], 1)
        self.assertEqual(options["measurement_counts"]["IVL"], 1)

    def test_spectral_calibration_creates_separate_workbook_and_model(self) -> None:
        source = self._seed_spectrum()
        before = source.read_bytes()
        options = self.service.options(self.series_path)
        group = next(group for group in options["groups"] if group["candidates"])

        started = self.service.start_calibration(self.series_path, {
            "selections": {
                group["key"]: {"pixel_id": self.pixel_id, "strategy": "replace"}
            },
            "thresholds": {
                "median_tolerance_percent": 12.0,
                "linear_model_outlier_percent": 55.0,
            },
        })
        self.assertTrue(started["active"])
        state = self._wait()

        self.assertEqual(state["status"], "completed", state)
        item = state["result"]["items"][0]
        self.assertTrue(Path(item["output"]).is_file())
        self.assertEqual(source.read_bytes(), before)
        manager = SeriesManager(self.series_path)
        self.assertIsNotNone(manager.integral_calibration_for_pixel(self.pixel_id))
        self.assertEqual(self.settings["spectral_calibration"]["median_tolerance_percent"], 12.0)

    def test_luminance_recalculation_requires_explicit_confirmation(self) -> None:
        self._seed_ivl()
        with self.assertRaises(RecalculationValidationError):
            self.service.start_luminance(self.series_path, {"confirmed": False})

    def test_luminance_recalculation_replaces_workbook_and_restores_raw(self) -> None:
        output = self._seed_ivl()
        started = self.service.start_luminance(self.series_path, {"confirmed": True})
        self.assertTrue(started["active"])

        state = self._wait()

        self.assertEqual(state["status"], "completed", state)
        self.assertEqual(state["result"]["workbooks_updated"], 1)
        self.assertEqual(state["result"]["raw_files_restored"], 1)
        workbook = load_workbook(output, read_only=True, data_only=True)
        try:
            self.assertIn("Cycle_1", workbook.sheetnames)
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()
