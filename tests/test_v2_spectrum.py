import tempfile
import time
import unittest
from pathlib import Path
import json
import urllib.error
import urllib.request

from openpyxl import load_workbook

from oled_app.measurements.ivl import IVLParams
from oled_app.settings import load_app_settings
from oled_v2.series_service import SeriesService, SeriesValidationError
from oled_v2.server import LocalBackend
from oled_v2.config import CLIENT_HEADER, SESSION_HEADER
from oled_v2.spectrum import SpectrumController


def wait_terminal(controller, timeout=15):
    deadline = time.monotonic() + timeout
    while controller.snapshot()["active"] and time.monotonic() < deadline:
        time.sleep(0.01)
    state = controller.snapshot()
    if state["active"]:
        raise AssertionError("Spectrum did not terminate")
    return state


FAST_PARAMS = {
    "voltage_start": 3.0,
    "voltage_end": 3.1,
    "voltage_step": 0.1,
    "settle_time_voltage_s": 0.0,
    "settle_time_spectrum_s": 0.0,
    "discard_first_scan_after_tint_change": False,
}


class SpectrumControllerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.controller = SpectrumController(self.root / "standalone")

    def tearDown(self):
        self.controller.shutdown()
        self.temp.cleanup()

    def create_series_with_ivl(self, keyword="queue"):
        service = SeriesService(self.root / keyword)
        active = service.create_series({
            "root": str(self.root / keyword), "deposition_date": "2026-09-27",
            "keyword": keyword, "series_led_color": "green",
            "quarter_bases": {str(number): "Q" for number in range(1, 5)},
            "quarter_descriptions": {str(number): "Test" for number in range(1, 5)},
        })["active"]
        for pixel in active["pixels"][:8]:
            target = {"series_path": active["path"], "pixel_id": pixel["pixel_id"]}
            context = service.ivl_target(target, IVLParams(), load_app_settings())
            service.record_ivl(context, IVLParams(), {
                "status": "WORKING", "file": str(Path(active["path"]) / f"{pixel['pixel_id']}.xlsx"),
                "run_id": "seed", "ivl_diagnosis": "seed", "opening_voltage": 3.0,
                "max_current_mA": 1.0, "max_photo_uA": 1.0,
            })
        return service, service.open_series(active["path"])["active"]

    def wait_decision(self, kind, timeout=12):
        deadline = time.monotonic() + timeout
        while time.monotonic() < deadline:
            state = self.controller.snapshot()
            if (state.get("decision") or {}).get("kind") == kind:
                return state
            if not state["active"]:
                raise AssertionError(f"Spectrum stopped before {kind}: {state}")
            time.sleep(0.01)
        raise AssertionError(f"Spectrum did not request {kind}")

    def decide(self, state, action, pixel_id=None):
        return self.controller.decide({
            "run_id": state["run_id"], "decision_id": state["decision"]["id"],
            "action": action, "pixel_id": pixel_id,
        })

    def test_invalid_preflight_is_read_only(self):
        for payload in (
            {"voltage_step": 0}, {"voltage_end": 1.0, "voltage_start": 2.0},
            {"target_intensity": 1000, "intensity_min": 2000},
            {"t_int_initial_s": 2, "t_int_max_s": 1},
            {"max_iterations": 1.5}, {"led_type": "ultraviolet"},
            {"hardware_mode": "real"},
        ):
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                self.controller.preflight(payload)
        self.assertFalse(self.controller.output_root.exists())

    def test_standalone_builds_compatible_workbook_and_preview(self):
        self.controller.start(FAST_PARAMS)
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertTrue(state["safe_shutdown_confirmed"])
        self.assertEqual(state["point_count"], 2)
        self.assertIsNotNone(state["latest_spectrum"])
        self.assertGreater(len(state["latest_spectrum"]["wavelengths_nm"]), 100)
        workbook = load_workbook(state["result"]["file"], read_only=True)
        try:
            self.assertIn("Сводка", workbook.sheetnames)
            self.assertGreaterEqual(len(workbook.sheetnames), 2)
        finally:
            workbook.close()
        state["points"].clear()
        self.assertEqual(len(self.controller.snapshot()["points"]), 2)

    def test_selected_series_pixel_is_journaled(self):
        service = SeriesService(self.root / "series")
        active = service.create_series({
            "root": str(self.root / "series"), "deposition_date": "2026-09-27",
            "keyword": "spectrum", "series_led_color": "green",
            "quarter_bases": {str(number): "Q" for number in range(1, 5)},
            "quarter_descriptions": {str(number): "Test" for number in range(1, 5)},
        })["active"]
        self.controller.series_service = service
        pixel_id = active["pixels"][0]["pixel_id"]
        self.controller.start({
            **FAST_PARAMS,
            "target": {"series_path": active["path"], "pixel_id": pixel_id},
            "use_opening_voltage": False,
        })
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertTrue(state["result"]["journaled"])
        reopened = service.open_series(active["path"])["active"]
        pixel = next(item for item in reopened["pixels"] if item["pixel_id"] == pixel_id)
        self.assertTrue(pixel["last_spectrum_file"])
        self.assertEqual(reopened["metrics"]["spectra"], 1)
        self.assertEqual(reopened["history"][-1]["type"], "SPECTRUM")
        self.assertIn("ЭМУЛЯТОР v2", reopened["history"][-1]["notes"])

    def test_opening_voltage_is_required_only_when_requested(self):
        service = SeriesService(self.root / "series-opening")
        active = service.create_series({
            "root": str(self.root / "series-opening"), "deposition_date": "2026-09-27",
            "keyword": "opening", "series_led_color": "blue",
            "quarter_bases": {str(number): "B" for number in range(1, 5)},
            "quarter_descriptions": {str(number): "Test" for number in range(1, 5)},
        })["active"]
        self.controller.series_service = service
        target = {"series_path": active["path"], "pixel_id": active["pixels"][0]["pixel_id"]}
        with self.assertRaisesRegex(SeriesValidationError, "нет напряжения открытия"):
            self.controller.preflight({**FAST_PARAMS, "target": target, "use_opening_voltage": True})
        checked = self.controller.preflight({**FAST_PARAMS, "target": target, "use_opening_voltage": False})
        self.assertEqual(checked["effective_voltage_start"], 3.0)

    def test_api_auth_busy_gate_and_snapshot_recovery(self):
        with LocalBackend(series_root=self.root / "api-series") as backend:
            controller = backend.server.config.app.state.spectrum_controller
            controller.output_root = self.root / "api-output"
            headers = {}

            def request(path, payload=None, method=None):
                req = urllib.request.Request(
                    backend.session.origin + path,
                    headers=headers,
                    data=json.dumps(payload).encode() if payload is not None else None,
                    method=method,
                )
                try:
                    with urllib.request.urlopen(req, timeout=5) as response:
                        return response.status, json.load(response)
                except urllib.error.HTTPError as exc:
                    return exc.code, json.load(exc)

            self.assertEqual(request("/api/spectrum/state")[0], 401)
            headers.update({
                SESSION_HEADER: backend.session.token,
                CLIENT_HEADER: "spectrum-test-client-0001",
                "Content-Type": "application/json",
            })
            self.assertEqual(request("/api/spectrum/preflight", {"voltage_step": 0})[0], 422)
            code, started = request("/api/spectrum/start", {
                "settle_time_voltage_s": 0.5,
                "settle_time_spectrum_s": 0.0,
                "discard_first_scan_after_tint_change": False,
            })
            self.assertEqual(code, 202)
            self.assertEqual(request("/api/ivl/start", {})[0], 409)
            self.assertEqual(request("/api/poc/start", {})[0], 409)
            self.assertEqual(request("/api/series/close", {}, "POST")[0], 409)
            snapshot = request("/api/spectrum/state")[1]
            self.assertEqual(snapshot["run_id"], started["run_id"])
            request("/api/spectrum/stop", {})
            state = wait_terminal(controller)
            self.assertEqual(state["status"], "stopped")
            self.assertTrue(state["safe_shutdown_confirmed"])

    def test_marked_substrate_queue_waits_for_each_pixel_and_persists_progress(self):
        service, active = self.create_series_with_ivl("marked-queue")
        self.controller.series_service = service
        selected = active["pixels"][:2]
        for pixel in selected:
            service.set_spectrum_priority(pixel["pixel_id"], True)
        queue = {
            "series_path": active["path"], "start_pixel": selected[0]["pixel_id"],
            "scope": "substrate", "queued_only": True,
        }
        checked = self.controller.preflight({**FAST_PARAMS, "queue": queue})
        self.assertEqual(checked["queue"]["total"], 2)
        priority = service.spectrum_queue_targets({
            "series_path": active["path"], "start_pixel": selected[1]["pixel_id"],
            "scope": "priority", "queued_only": True,
        })
        self.assertEqual([item["pixel_id"] for item in priority["targets"]], [selected[1]["pixel_id"]])
        self.controller.start({**FAST_PARAMS, "queue": queue})
        for expected in selected:
            waiting = self.wait_decision("next_pixel")
            self.assertEqual(waiting["decision"]["pixel_id"], expected["pixel_id"])
            self.decide(waiting, "measure")
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertEqual(state["queue"]["completed"], 2)
        self.assertEqual(state["queue"]["remaining"], 0)
        self.assertEqual(state["queue"]["attempts"], 2)
        reopened = service.open_series(active["path"])["active"]
        self.assertEqual(reopened["metrics"]["spectra"], 2)
        self.assertEqual(reopened["metrics"]["spectrum_queue"], 0)

    def test_no_contact_queue_can_retry_same_pixel_then_continue(self):
        service, active = self.create_series_with_ivl("no-contact-queue")
        self.controller.series_service = service
        pixel = active["pixels"][3]
        service.set_spectrum_priority(pixel["pixel_id"], True)
        queue = {
            "series_path": active["path"], "start_pixel": pixel["pixel_id"],
            "scope": "substrate", "queued_only": True,
        }
        self.controller.start({**FAST_PARAMS, "queue": queue})
        self.decide(self.wait_decision("next_pixel"), "measure")
        first = self.wait_decision("no_contact")
        self.assertTrue(first["safe_shutdown_confirmed"])
        self.decide(first, "retry")
        second = self.wait_decision("no_contact")
        self.decide(second, "continue")
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed")
        self.assertEqual(state["queue"]["attempts"], 2)
        self.assertEqual([item["pixel_id"] for item in state["queue"]["results"]], [pixel["pixel_id"]] * 2)
        self.assertTrue(all(item["file"] is None for item in state["queue"]["results"]))

    def test_current_limit_queue_can_delete_data_and_offer_same_quarter_replacement(self):
        service, active = self.create_series_with_ivl("limit-queue")
        self.controller.series_service = service
        pixel = active["pixels"][4]
        service.set_spectrum_priority(pixel["pixel_id"], True)
        queue = {
            "series_path": active["path"], "start_pixel": pixel["pixel_id"],
            "scope": "substrate", "queued_only": True,
        }
        params = {
            **FAST_PARAMS, "voltage_start": 3.5, "voltage_end": 3.5,
            "queue": queue, "use_opening_voltage": False,
        }
        self.controller.start(params)
        self.decide(self.wait_decision("next_pixel"), "measure")
        rejected = self.wait_decision("rejected_data")
        self.assertEqual(rejected["decision"]["status"], "CURRENT_LIMIT")
        self.assertTrue(rejected["safe_shutdown_confirmed"])
        self.decide(rejected, "delete")
        replacement = self.wait_decision("replacement")
        self.assertTrue(replacement["decision"]["replacement_pixels"])
        replacement_id = replacement["decision"]["replacement_pixels"][0]
        self.decide(replacement, "replace", replacement_id)
        next_pixel = self.wait_decision("next_pixel")
        self.assertEqual(next_pixel["decision"]["pixel_id"], replacement_id)
        self.decide(next_pixel, "skip")
        state = wait_terminal(self.controller)
        self.assertEqual(state["status"], "completed", state)
        self.assertEqual(state["queue"]["attempts"], 1)
        self.assertEqual(state["queue"]["skipped_pixels"], [replacement_id])
        self.assertIsNone(state["queue"]["results"][0]["file"])
        reopened = service.open_series(active["path"])["active"]
        saved = next(item for item in reopened["pixels"] if item["pixel_id"] == pixel["pixel_id"])
        self.assertEqual(saved["status"], "CURRENT_LIMIT")
        self.assertIsNone(saved["last_spectrum_file"])

    def test_stale_queue_decision_is_rejected(self):
        service, active = self.create_series_with_ivl("stale-decision")
        self.controller.series_service = service
        pixel = active["pixels"][0]
        service.set_spectrum_priority(pixel["pixel_id"], True)
        self.controller.start({
            **FAST_PARAMS,
            "queue": {"series_path": active["path"], "start_pixel": pixel["pixel_id"],
                      "scope": "substrate", "queued_only": True},
        })
        waiting = self.wait_decision("next_pixel")
        self.decide(waiting, "skip")
        with self.assertRaises(RuntimeError):
            self.decide(waiting, "skip")
        state = wait_terminal(self.controller)
        self.assertEqual(state["queue"]["skipped_pixels"], [pixel["pixel_id"]])


if __name__ == "__main__":
    unittest.main()
