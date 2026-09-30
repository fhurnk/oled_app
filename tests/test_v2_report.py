from __future__ import annotations

import json
from pathlib import Path
import tempfile
import time
import unittest

from openpyxl import Workbook, load_workbook

from oled_v2.report import ReportService, ReportValidationError


def _write_spectrum(path: Path, pixel: str, voltages=(2.0, 2.1)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Pixel", pixel])
    processed = workbook.create_sheet("Processed counts per s")
    processed.append(["V set (V)", *voltages])
    processed.append(["Wavelength (nm)", *["Processed counts per s"] * len(voltages)])
    processed.append([450.0, *[1000.0 + index for index in range(len(voltages))]])
    processed.append([550.0, *[2000.0 + index for index in range(len(voltages))]])
    workbook.save(path)


class V2ReportServiceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.series = Path(self.temp.name) / "series"
        self.series.mkdir()
        (self.series / "series_config.json").write_text(
            json.dumps({
                "description_scope": "half",
                "half_orientation": "top_bottom",
                "quarter_bases": {"1": "CG", "2": "CG", "3": "CG", "4": "CG"},
                "quarter_descriptions": {"1": "top", "2": "top", "3": "bottom", "4": "bottom"},
            }),
            encoding="utf-8",
        )
        self.service = ReportService()

    def tearDown(self) -> None:
        self.service.shutdown()
        self.temp.cleanup()

    def _wait(self) -> dict:
        deadline = time.monotonic() + 8.0
        while time.monotonic() < deadline:
            state = self.service.snapshot()
            if not state["active"]:
                return state
            time.sleep(0.02)
        self.fail("Report did not finish in time")

    def test_options_group_spectra_by_series_scope(self) -> None:
        root = self.series / "measurements" / "02_SPECTRA" / "2026-09-30"
        _write_spectrum(root / "CG1" / "CG1_1" / "CG1_1_1" / "SPECTRUM_CG1_1_1.xlsx", "CG1_1_1")
        _write_spectrum(root / "CG2" / "CG2_1" / "CG2_1_1" / "SPECTRUM_CG2_1_1.xlsx", "CG2_1_1")

        options = self.service.options(self.series, {"grouping": "settings"})

        self.assertEqual(options["available_modes"], ["spectra"])
        self.assertEqual([group["key"] for group in options["groups"]], ["2+1"])
        pixels = {
            pixel["pixel_id"]
            for substrate in options["groups"][0]["substrates"]
            for pixel in substrate["pixels"]
        }
        self.assertEqual(pixels, {"CG1_1_1", "CG2_1_1"})

    def test_ivl_only_xlsx_runs_in_background(self) -> None:
        (self.series / "measurements" / "01_IVL_VAH" / "2026-09-30").mkdir(parents=True)

        started = self.service.start(self.series, {
            "mode": "ivl", "grouping": "settings", "ivl_date": "2026-09-30",
            "spectrum_date": "", "excluded_quarters": [], "format": "xlsx",
            "output_name": "report_test.xlsx",
        })
        self.assertTrue(started["active"])
        completed = self._wait()

        self.assertEqual(completed["status"], "completed", completed["error"])
        output = self.series / "report_test.xlsx"
        self.assertTrue(output.is_file())
        workbook = load_workbook(output, read_only=True)
        try:
            self.assertIn("IVL_U_I_PD", workbook.sheetnames)
            self.assertNotIn("Spectra_by_voltage", workbook.sheetnames)
        finally:
            workbook.close()

    def test_spectrum_xlsx_uses_explicit_pixel_and_grid(self) -> None:
        root = self.series / "measurements" / "02_SPECTRA" / "2026-09-30"
        _write_spectrum(root / "CG1" / "CG1_1" / "CG1_1_1" / "SPECTRUM_CG1_1_1.xlsx", "CG1_1_1")

        self.service.start(self.series, {
            "mode": "spectra", "grouping": "quarters", "ivl_date": "",
            "spectrum_date": "2026-09-30", "excluded_quarters": [], "format": "xlsx",
            "output_name": "spectrum_report.xlsx",
            "selection": {"1": {"substrate": "CG1_1", "pixel_id": "CG1_1_1"}},
            "same_grid": True,
            "global_grid": {"start": 2.0, "stop": 2.1, "step": 0.1},
            "pixel_grids": {},
        })
        completed = self._wait()

        self.assertEqual(completed["status"], "completed", completed["error"])
        self.assertEqual(completed["result"]["spectrum_records"], 1)
        self.assertTrue((self.series / "spectrum_report.xlsx").is_file())

    def test_output_cannot_escape_or_overwrite_series(self) -> None:
        (self.series / "measurements" / "01_IVL_VAH" / "2026-09-30").mkdir(parents=True)
        common = {
            "mode": "ivl", "grouping": "settings", "ivl_date": "2026-09-30",
            "spectrum_date": "", "excluded_quarters": [], "format": "xlsx",
        }
        with self.assertRaisesRegex(ReportValidationError, "без папок"):
            self.service.start(self.series, {**common, "output_name": "..\\outside.xlsx"})
        (self.series / "exists.xlsx").write_bytes(b"keep")
        with self.assertRaisesRegex(ReportValidationError, "уже существует"):
            self.service.start(self.series, {**common, "output_name": "exists.xlsx"})


if __name__ == "__main__":
    unittest.main()
