"""Backend-owned report discovery and generation for the v2 desktop shell."""

from __future__ import annotations

from copy import deepcopy
from datetime import datetime, timezone
import math
from pathlib import Path
import threading
from typing import Any, Dict, Optional

from openpyxl import load_workbook

from oled_app.constants import MEASUREMENT_FOLDER_NAMES
from oled_app.reports.origin_report import (
    REPORT_GROUPING_QUARTERS,
    REPORT_GROUPING_SETTINGS,
    REPORT_GROUPINGS,
    REPORT_MODE_FULL,
    REPORT_MODE_IVL,
    REPORT_MODE_SPECTRA,
    REPORT_MODES,
    build_origin_project,
    build_workbook,
    collect_report_data,
    load_series_config,
    parse_args,
    report_group_key,
    report_group_label,
    series_quarter_number,
)
from oled_app.utils import build_report_voltage_grid, voltage_grid_missing


class ReportError(RuntimeError):
    pass


class ReportValidationError(ReportError):
    pass


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _measurement_dates(series_path: Path, measurement_type: str) -> list[str]:
    folder = series_path / "measurements" / MEASUREMENT_FOLDER_NAMES[measurement_type]
    if not folder.exists():
        return []
    return sorted(path.name for path in folder.iterdir() if path.is_dir())


def _spectrum_voltages(path: Path) -> list[float]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        if "Processed counts per s" not in workbook.sheetnames:
            return []
        sheet = workbook["Processed counts per s"]
        voltage_row = None
        for row in range(1, min(sheet.max_row, 30) + 1):
            if sheet.cell(row=row, column=1).value == "V set (V)":
                voltage_row = row
                break
        if voltage_row is None:
            return []
        return [
            round(float(value), 6)
            for column in range(2, sheet.max_column + 1)
            for value in [sheet.cell(row=voltage_row, column=column).value]
            if isinstance(value, (int, float))
        ]
    finally:
        workbook.close()


def _default_step(voltages: list[float]) -> float:
    values = sorted(set(round(float(value), 6) for value in voltages))
    differences = [round(right - left, 6) for left, right in zip(values, values[1:]) if right > left]
    return min(differences) if differences else 0.1


def _output_name(ivl_date: str, spectrum_date: str, mode: str, suffix: str) -> str:
    if mode == REPORT_MODE_IVL:
        return f"report_IVL_{ivl_date}{suffix}"
    if mode == REPORT_MODE_SPECTRA:
        return f"report_Spctr_{spectrum_date}{suffix}"
    stem = ivl_date if ivl_date == spectrum_date else f"IVL_{ivl_date}_Spctr_{spectrum_date}"
    return f"report_{stem}{suffix}"


class ReportService:
    def __init__(self, logger=None) -> None:
        self._logger = logger
        self._lock = threading.RLock()
        self._thread: Optional[threading.Thread] = None
        self._state: Dict[str, Any] = {
            "status": "idle", "active": False, "started_at": None,
            "finished_at": None, "error": None, "result": None,
            "series_path": None,
        }

    def snapshot(self) -> Dict[str, Any]:
        with self._lock:
            return deepcopy(self._state)

    def state(self, series_path: Optional[Path], payload: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        generation = self.snapshot()
        if series_path is None:
            return {"available": False, "generation": generation, "options": None}
        if (generation.get("series_path")
                and generation["series_path"] != str(Path(series_path).resolve())
                and not generation["active"]):
            generation = {
                "status": "idle", "active": False, "started_at": None,
                "finished_at": None, "error": None, "result": None,
                "series_path": None,
            }
        return {
            "available": True,
            "generation": generation,
            "options": self.options(Path(series_path), payload or {}),
        }

    def options(self, series_path: Path, payload: Dict[str, Any]) -> Dict[str, Any]:
        grouping = str(payload.get("grouping") or REPORT_GROUPING_SETTINGS)
        if grouping not in REPORT_GROUPINGS:
            raise ReportValidationError("Неизвестный режим группировки отчёта.")
        excluded = self._excluded(payload.get("excluded_quarters", []))
        ivl_dates = _measurement_dates(series_path, "IVL")
        spectrum_dates = _measurement_dates(series_path, "SPECTRUM")
        spectrum_date = str(payload.get("spectrum_date") or (spectrum_dates[-1] if spectrum_dates else ""))
        if spectrum_date and spectrum_date not in spectrum_dates:
            raise ReportValidationError("Выбранная дата спектров отсутствует в серии.")
        candidates = self._candidates(series_path, spectrum_date, excluded, grouping)
        modes = []
        if ivl_dates and spectrum_dates:
            modes.append(REPORT_MODE_FULL)
        if ivl_dates:
            modes.append(REPORT_MODE_IVL)
        if spectrum_dates:
            modes.append(REPORT_MODE_SPECTRA)
        config = load_series_config(series_path / "measurements")
        groups = []
        for group_key in sorted(candidates):
            substrates = []
            for substrate in sorted(candidates[group_key]):
                pixels = [
                    {"pixel_id": pixel, "voltages": info["voltages"]}
                    for pixel, info in sorted(candidates[group_key][substrate].items())
                ]
                substrates.append({"name": substrate, "pixels": pixels})
            groups.append({
                "key": group_key,
                "label": report_group_label(config, group_key),
                "substrates": substrates,
            })
        return {
            "series_path": str(series_path.resolve()),
            "ivl_dates": ivl_dates,
            "spectrum_dates": spectrum_dates,
            "available_modes": modes,
            "groups": groups,
        }

    def start(self, series_path: Path, payload: Dict[str, Any]) -> Dict[str, Any]:
        if not isinstance(payload, dict):
            raise ReportValidationError("Параметры отчёта должны быть объектом.")
        with self._lock:
            if self._state["active"]:
                raise ReportValidationError("Отчёт уже создаётся.")
        args, output = self._build_args(Path(series_path), payload)
        with self._lock:
            self._state = {
                "status": "running", "active": True, "started_at": _utc_now(),
                "finished_at": None, "error": None, "result": None,
                "series_path": str(Path(series_path).resolve()),
            }
            self._thread = threading.Thread(
                target=self._run,
                args=(args, output),
                daemon=True,
                name="oled-v2-report",
            )
            self._thread.start()
            return self.snapshot()

    def shutdown(self) -> None:
        thread = self._thread
        if thread is not None and thread.is_alive():
            thread.join(timeout=2.0)

    def _run(self, args, output: Path) -> None:
        try:
            data = collect_report_data(args)
            if args.format == "origin":
                warnings = build_origin_project(args, data)
            else:
                workbook, warnings = build_workbook(args, data)
                output.parent.mkdir(parents=True, exist_ok=True)
                workbook.save(output)
            result = {
                "output": str(output),
                "format": args.format,
                "ivl_records": len(data.iv_records),
                "spectrum_records": len(data.spectrum_records),
                "warnings": list(warnings),
            }
            with self._lock:
                self._state.update({
                    "status": "completed", "active": False, "finished_at": _utc_now(),
                    "result": result,
                })
            if self._logger:
                self._logger.info("Report created: %s", output)
        except Exception as exc:
            with self._lock:
                self._state.update({
                    "status": "failed", "active": False, "finished_at": _utc_now(),
                    "error": str(exc),
                })
            if self._logger:
                self._logger.exception("Report failed")

    def _build_args(self, series_path: Path, payload: Dict[str, Any]):
        mode = str(payload.get("mode") or REPORT_MODE_FULL)
        grouping = str(payload.get("grouping") or REPORT_GROUPING_SETTINGS)
        output_format = str(payload.get("format") or "xlsx")
        if mode not in REPORT_MODES:
            raise ReportValidationError("Неизвестный состав отчёта.")
        if grouping not in REPORT_GROUPINGS:
            raise ReportValidationError("Неизвестная группировка отчёта.")
        if output_format not in {"origin", "xlsx"}:
            raise ReportValidationError("Формат отчёта должен быть origin или xlsx.")
        excluded = self._excluded(payload.get("excluded_quarters", []))
        options = self.options(series_path, {
            "grouping": grouping,
            "excluded_quarters": sorted(excluded),
            "spectrum_date": payload.get("spectrum_date"),
        })
        if mode not in options["available_modes"]:
            raise ReportValidationError("Для выбранного состава в серии недостаточно данных.")
        ivl_date = str(payload.get("ivl_date") or (options["ivl_dates"][-1] if options["ivl_dates"] else ""))
        spectrum_date = str(payload.get("spectrum_date") or (options["spectrum_dates"][-1] if options["spectrum_dates"] else ""))
        if mode in {REPORT_MODE_FULL, REPORT_MODE_IVL} and ivl_date not in options["ivl_dates"]:
            raise ReportValidationError("Выбранная дата ВАЯХ отсутствует в серии.")

        suffix = ".opju" if output_format == "origin" else ".xlsx"
        default_name = _output_name(ivl_date, spectrum_date, mode, suffix)
        output_name = str(payload.get("output_name") or default_name).strip()
        if Path(output_name).name != output_name or Path(output_name).suffix.lower() != suffix:
            raise ReportValidationError(f"Имя файла должно быть без папок и оканчиваться на {suffix}.")
        output = (series_path / output_name).resolve()
        if output.parent != series_path.resolve():
            raise ReportValidationError("Отчёт можно сохранить только в папке активной серии.")
        if output.exists():
            raise ReportValidationError("Файл с таким именем уже существует. Задайте другое имя.")

        argv = [
            "--measurements-dir", str(series_path / "measurements"),
            "--output", str(output), "--format", output_format,
            "--report-mode", mode, "--report-grouping", grouping, "--strict",
        ]
        for quarter in sorted(excluded):
            argv.extend(["--exclude-quarter", str(quarter)])
        if mode in {REPORT_MODE_FULL, REPORT_MODE_IVL}:
            argv.extend(["--ivl-date", ivl_date])
        if mode in {REPORT_MODE_FULL, REPORT_MODE_SPECTRA}:
            argv.extend(["--spectrum-date", spectrum_date, "--require-spectrum-pixel-selection"])
            selected = self._selected(options["groups"], payload.get("selection"))
            for group, info in sorted(selected.items()):
                argv.extend(["--spectrum-group-pixel", f"{group}={info['pixel_id']}"])
            same_grid = payload.get("same_grid", True)
            if not isinstance(same_grid, bool):
                raise ReportValidationError("Режим сетки напряжений должен быть логическим.")
            if same_grid:
                grid = self._grid(payload.get("global_grid"), "Общая сетка")
                requested = build_report_voltage_grid(*grid)
                for info in selected.values():
                    missing = voltage_grid_missing(requested, info["voltages"])
                    if missing:
                        raise ReportValidationError(
                            f"{info['pixel_id']}: отсутствуют напряжения {', '.join(map(str, missing[:8]))}."
                        )
                argv.extend(["--voltage-start", str(grid[0]), "--voltage-stop", str(grid[1]), "--voltage-step", str(grid[2])])
            else:
                grids = payload.get("pixel_grids")
                if not isinstance(grids, dict):
                    raise ReportValidationError("Задайте индивидуальные сетки спектров.")
                common: Optional[set[float]] = None
                for info in selected.values():
                    pixel = info["pixel_id"]
                    grid = self._grid(grids.get(pixel), f"Сетка {pixel}")
                    requested = build_report_voltage_grid(*grid)
                    missing = voltage_grid_missing(requested, info["voltages"])
                    if missing:
                        raise ReportValidationError(f"{pixel}: отсутствуют напряжения {', '.join(map(str, missing[:8]))}.")
                    values = set(round(item, 6) for item in requested)
                    common = values if common is None else common & values
                    argv.extend(["--spectrum-voltage-grid", f"{pixel}={grid[0]}:{grid[1]}:{grid[2]}"])
                if not common:
                    raise ReportValidationError("У индивидуальных сеток нет общего напряжения.")
        return parse_args(argv), output

    @staticmethod
    def _excluded(value: Any) -> set[int]:
        if not isinstance(value, list):
            raise ReportValidationError("Исключаемые четверти должны быть списком.")
        result = set()
        for item in value:
            if isinstance(item, bool) or not isinstance(item, int) or item not in range(1, 5):
                raise ReportValidationError("Номер исключаемой четверти должен быть от 1 до 4.")
            result.add(item)
        return result

    @staticmethod
    def _grid(value: Any, label: str) -> tuple[float, float, float]:
        if not isinstance(value, dict):
            raise ReportValidationError(f"{label}: задайте начало, конец и шаг.")
        numbers = []
        for key in ("start", "stop", "step"):
            item = value.get(key)
            if isinstance(item, bool) or not isinstance(item, (int, float)):
                raise ReportValidationError(f"{label}: {key} должен быть числом.")
            number = float(item)
            if not math.isfinite(number):
                raise ReportValidationError(f"{label}: {key} должен быть конечным числом.")
            numbers.append(number)
        start, stop, step = numbers
        if step <= 0 or stop < start:
            raise ReportValidationError(f"{label}: требуется начало ≤ конец и шаг > 0.")
        if len(build_report_voltage_grid(start, stop, step)) > 1000:
            raise ReportValidationError(f"{label}: допускается не более 1000 точек.")
        return start, stop, step

    @staticmethod
    def _selected(groups: list[dict], value: Any) -> Dict[str, Dict[str, Any]]:
        if not isinstance(value, dict):
            raise ReportValidationError("Выберите подложку и пиксель для каждой группы.")
        selected = {}
        for group in groups:
            choice = value.get(group["key"])
            if not isinstance(choice, dict):
                raise ReportValidationError(f"Не выбран спектр для группы {group['label']}.")
            substrate_name = str(choice.get("substrate") or "")
            pixel_id = str(choice.get("pixel_id") or "")
            match = None
            for substrate in group["substrates"]:
                if substrate["name"] != substrate_name:
                    continue
                match = next((pixel for pixel in substrate["pixels"] if pixel["pixel_id"] == pixel_id), None)
            if match is None:
                raise ReportValidationError(f"Выбранный спектр группы {group['label']} не найден.")
            selected[group["key"]] = match
        if not selected:
            raise ReportValidationError("Не найдено спектров для отчёта.")
        return selected

    @staticmethod
    def _candidates(
        series_path: Path,
        spectrum_date: str,
        excluded: set[int],
        grouping: str,
    ) -> Dict[str, Dict[str, Dict[str, Dict[str, Any]]]]:
        root = series_path / "measurements" / MEASUREMENT_FOLDER_NAMES["SPECTRUM"]
        if not root.exists() or not spectrum_date:
            return {}
        config = load_series_config(series_path / "measurements")
        latest: Dict[str, tuple[float, str, str, Path]] = {}
        for path in root.rglob("SPECTRUM_*.xlsx"):
            try:
                parts = path.relative_to(root).parts
            except ValueError:
                continue
            if len(parts) < 5 or parts[0] != spectrum_date:
                continue
            series, substrate, pixel = parts[1], parts[2], parts[3]
            quarter = series_quarter_number(series)
            if quarter in excluded:
                continue
            previous = latest.get(pixel)
            modified = path.stat().st_mtime
            if previous is None or modified > previous[0]:
                latest[pixel] = (modified, series, substrate, path)
        candidates: Dict[str, Dict[str, Dict[str, Dict[str, Any]]]] = {}
        for pixel, (_modified, series, substrate, path) in latest.items():
            voltages = _spectrum_voltages(path)
            quarter = series_quarter_number(series)
            if not voltages or quarter is None:
                continue
            group = report_group_key(config, quarter, grouping)
            candidates.setdefault(group, {}).setdefault(substrate, {})[pixel] = {
                "voltages": voltages,
                "step": _default_step(voltages),
            }
        return candidates
